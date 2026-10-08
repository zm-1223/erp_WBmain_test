# -*- coding: utf-8 -*-
"""从 test.md 同步两张表到仓库与桌面 xlsx。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "test.md"
OUTS = [
    ROOT / "erp_WB_test.xlsx",
    Path.home() / "Desktop" / "erp_WB_test.xlsx",
]


def _tables(md: str) -> list[list[list[str]]]:
    tables = []
    cur = []
    for line in md.splitlines():
        if line.startswith("| 序号"):
            if cur:
                tables.append(cur)
            cur = [line]
        elif cur and line.startswith("|"):
            if set(line.replace("|", "").strip()) <= set("- "):
                continue
            cur.append(line)
        elif cur and not line.startswith("|"):
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    parsed = []
    for t in tables:
        rows = []
        for line in t:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            rows.append(cells)
        parsed.append(rows)
    return parsed


def main():
    tables = _tables(MD.read_text(encoding="utf-8"))
    wb = Workbook()
    names = ["WB刊登", "Ozon刊登"]
    for i, rows in enumerate(tables[:2]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = names[i] if i < len(names) else f"Sheet{i+1}"
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["G"].width = 50
        ws.column_dimensions["H"].width = 40
    for dest in OUTS:
        try:
            dest.parent.mkdir(parents=True, exist_ok=True)
            wb.save(dest)
            print("wrote", dest)
        except Exception as e:
            print("skip", dest, e)


if __name__ == "__main__":
    main()
